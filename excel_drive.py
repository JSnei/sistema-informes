import os
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseUpload

from database.database import obtener_servicios_para_excel


# ==========================================================
# CONFIGURACIÓN GOOGLE DRIVE
# ==========================================================

GOOGLE_CREDENTIALS_FILE = os.getenv(
    "GOOGLE_DRIVE_CREDENTIALS",
    "/home/bsadmin/google-drive-credentials.json"
)

GOOGLE_DRIVE_FOLDER_ID = os.getenv(
    "GOOGLE_DRIVE_FOLDER_ID"
)

NOMBRE_ARCHIVO = "Servicios.xlsx"


# ==========================================================
# GENERAR EXCEL DESDE POSTGRESQL
# ==========================================================

def generar_excel_servicios():

    servicios = obtener_servicios_para_excel()

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Servicios"

    encabezados = [
        "N° Servicio",
        "N° Solicitud",
        "Fecha",
        "Ticket",
        "SAP Tienda",
        "Nombre Tienda",
        "Ciudad",
        "Departamento",
        "Dirección",
        "Región",
        "Tipo Tienda",
        "Tipo Servicio",
        "Cédula Técnico",
        "Nombre Técnico",
        "SAP Encargado",
        "Nombre Encargado",
        "Descripción Falla",
        "Trabajo Realizado",
        "Observaciones",
        "¿Suministró Repuestos?",
        "Repuestos Suministrados",
        "Estado",
    ]

    hoja.append(encabezados)

    for celda in hoja[1]:
        celda.font = Font(bold=True)
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    for servicio in servicios:

        hoja.append([
            servicio["servicio_id"],
            servicio["solicitud_id"],
            servicio["fecha"].replace(tzinfo=None) if servicio["fecha"] else None,
            servicio["ticket"],
            servicio["sap_tienda"],
            servicio["nombre_tienda"],
            servicio["ciudad"],
            servicio["departamento"],
            servicio["direccion"],
            servicio["region"],
            servicio["tipo_tienda"],
            servicio["tipo_servicio"],
            servicio["cedula_tecnico"],
            servicio["nombre_tecnico"],
            servicio["sap_encargado"],
            servicio["nombre_encargado"],
            servicio["descripcion_falla"],
            servicio["trabajo_realizado"],
            servicio["observaciones"],
            servicio["suministro_repuestos"],
            servicio["repuestos_suministrados"],
            servicio["estado"],
        ])

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    for columna in range(1, hoja.max_column + 1):

        letra = get_column_letter(columna)
        ancho_maximo = 0

        for celda in hoja[letra]:

            if celda.value is not None:
                ancho_maximo = max(
                    ancho_maximo,
                    len(str(celda.value))
                )

        hoja.column_dimensions[letra].width = min(
            ancho_maximo + 2,
            45
        )

    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)

    return archivo


# ==========================================================
# CONECTAR CON GOOGLE DRIVE
# ==========================================================

def obtener_drive():

    if not GOOGLE_DRIVE_FOLDER_ID:
        raise RuntimeError(
            "GOOGLE_DRIVE_FOLDER_ID no está configurado."
        )

    credenciales = (
        service_account.Credentials.from_service_account_file(
            GOOGLE_CREDENTIALS_FILE,
            scopes=[
                "https://www.googleapis.com/auth/drive"
            ]
        )
    )

    return build(
        "drive",
        "v3",
        credentials=credenciales,
        cache_discovery=False
    )


# ==========================================================
# SUBIR O ACTUALIZAR SERVICIOS.XLSX
# ==========================================================

def sincronizar_excel_drive():

    drive = obtener_drive()
    archivo = generar_excel_servicios()

    media = MediaIoBaseUpload(
        archivo,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        resumable=False
    )

    consulta = (
        f"name='{NOMBRE_ARCHIVO}' "
        f"and '{GOOGLE_DRIVE_FOLDER_ID}' in parents "
        "and trashed=false"
    )

    resultado = drive.files().list(
        q=consulta,
        fields="files(id, name)"
    ).execute()

    archivos = resultado.get("files", [])

    # Si ya existe, reemplazamos su contenido.
    if archivos:

        archivo_id = archivos[0]["id"]

        drive.files().update(
            fileId=archivo_id,
            media_body=media
        ).execute()

        return {
            "accion": "actualizado",
            "archivo_id": archivo_id
        }

    # Si todavía no existe, lo creamos.
    metadata = {
        "name": NOMBRE_ARCHIVO,
        "parents": [GOOGLE_DRIVE_FOLDER_ID]
    }

    creado = drive.files().create(
        body=metadata,
        media_body=media,
        fields="id"
    ).execute()

    return {
        "accion": "creado",
        "archivo_id": creado["id"]
    }


# ==========================================================
# SINCRONIZACIÓN SEGURA
# ==========================================================

def sincronizar_excel_seguro():

    try:
        resultado = sincronizar_excel_drive()

        print(
            f"Excel Drive sincronizado: "
            f"{resultado['accion']}"
        )

        return True

    except Exception as error:

        print(
            f"ADVERTENCIA: No se pudo sincronizar "
            f"Servicios.xlsx con Google Drive: {error}"
        )

        return False