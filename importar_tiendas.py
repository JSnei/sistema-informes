from openpyxl import load_workbook

from database.database import (
    crear_base_datos,
    agregar_tienda,
    obtener_tienda_por_sap
)


ARCHIVO_EXCEL = "Tiendas_Unificadas_R7_R8_CORREGIDO.xlsx"


def importar_tiendas():

    print("==========================================")
    print("IMPORTACIÓN DE TIENDAS")
    print("==========================================")

    crear_base_datos()

    libro = load_workbook(
        ARCHIVO_EXCEL,
        data_only=True
    )

    hoja = libro.active

    nuevas = 0
    existentes = 0
    errores = 0

    for fila in hoja.iter_rows(
        min_row=2,
        values_only=True
    ):

        sap = fila[0]
        region = fila[1]
        tipo = fila[2]
        nombre = fila[3]
        ciudad = fila[4]
        departamento = fila[5]
        direccion = fila[6]

        if sap is None:
            continue

        sap = str(sap).strip()

        try:

            tienda_existente = obtener_tienda_por_sap(sap)

            if tienda_existente:

                print(
                    f"YA EXISTE: "
                    f"{sap} - {tienda_existente.nombre}"
                )

                existentes += 1
                continue

            agregar_tienda(
                sap=sap,
                region=region,
                tipo=tipo,
                nombre=nombre,
                ciudad=ciudad,
                departamento=departamento,
                direccion=direccion
            )

            print(
                f"AGREGADA: "
                f"{sap} - {nombre}"
            )

            nuevas += 1

        except Exception as error:

            print(
                f"ERROR EN SAP {sap}: {error}"
            )

            errores += 1


    libro.close()


    print("")
    print("==========================================")
    print("IMPORTACIÓN FINALIZADA")
    print("==========================================")
    print(f"Tiendas nuevas:      {nuevas}")
    print(f"Tiendas existentes: {existentes}")
    print(f"Errores:             {errores}")
    print("==========================================")


if __name__ == "__main__":
    importar_tiendas()